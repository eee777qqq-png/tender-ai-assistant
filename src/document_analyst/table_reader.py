"""Извлечение таблиц (не плоского текста) из вложений документации закупки —
нужно там, где важна структура строки "наименование | ед.изм | количество",
которую `docx_reader.extract_text_from_docx()` намеренно теряет (тот модуль
просто склеивает весь текст в один поток для regex-поиска Агента 3).

Новый потребитель — `smeta_estimator.work_volume_extractor` (Агент 4, п.8
"Известных пробелов" в CLAUDE.md): ведомость объёмов работ реальной закупки
обычно живёт именно в такой таблице внутри приложения документации (.docx
или .xlsx), не в структурированных полях извещения.

**Честно не проверено на реальном документе.** В этой сессии нет доступа к
токену ЕИС и реальным скачанным архивам (см. CLAUDE.md, «Известные
пробелы», п.10 — они не переживают пересоздание облачного контейнера).
Формат таблиц ниже реконструирован по тому, как обычно оформляют ведомости
объёмов работ/локальные сметные расчёты в документации закупок на
капремонт/строительство (колонки "№ | Наименование работ | Ед. изм. |
Количество", иногда с доп. колонками цены) — не подтверждён на конкретном
файле этого проекта. Первый реальный документ такого типа может показать
вариант оформления, не предусмотренный здесь, как уже было с Агентом 3
(см. CLAUDE.md, открытый п.3) — это ожидаемо, не повод считать модуль
готовым к продакшену без проверки."""

from __future__ import annotations

import io

from docx import Document
from openpyxl import load_workbook

Row = list[str]
Table = list[Row]


def extract_docx_tables(content: bytes) -> list[Table]:
    """Каждая таблица документа — как список строк, каждая строка — список
    текстов ячеек, в порядке появления. Пустые документы/документы без
    таблиц — пустой список, не ошибка."""
    document = Document(io.BytesIO(content))
    tables: list[Table] = []
    for table in document.tables:
        rows: list[Row] = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)
    return tables


def extract_xlsx_tables(content: bytes) -> list[Table]:
    """Каждый лист книги — одна "таблица" (список строк с текстом ячеек,
    пустая ячейка -> ""). `data_only=True` — берём посчитанные значения
    формул, а не сами формулы, раз книга это позволяет (иначе ячейка с
    формулой без сохранённого кэша значения останется `None`/пустой —
    честно, не строим собственный движок формул)."""
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    tables: list[Table] = []
    for sheet in workbook.worksheets:
        rows: list[Row] = []
        for row in sheet.iter_rows():
            cells = ["" if cell.value is None else str(cell.value).strip() for cell in row]
            if any(cells):
                rows.append(cells)
        if rows:
            tables.append(rows)
    return tables
