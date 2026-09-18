"""Разбор данных из `regional_pricing_client.py` — текущих цен и индексов ГОСР.

Формат подтверждён на реальном файле для Москвы, 3 квартал 2026
(`tests/fixtures/gosr_moscow_q3_2026_sample.xlsx` — вырезанные настоящие
строки, не выдуманные), скачанном 2026-09-18.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from typing import Iterable

import openpyxl

from regulatory_updates.models import RegulatoryVersion

GOSR_SOURCE_ID_TEMPLATE = "fgiscs_gosr_index_{region_key}"


@dataclass
class GosrIndexEntry:
    """Одна строка из отчёта ГОСР — индекс пересчёта для конкретного кода
    ресурса, привязанный к его группе однородных строительных ресурсов."""

    resource_code: str
    resource_name: str
    unit: str
    base_price_2022: float
    group_number: int
    group_name: str
    index_value: float


def parse_current_prices_json(raw_json: bytes | str) -> dict[str, float]:
    """`GET .../BuildingResources/Search/{Materials|Machines}` группирует
    позиции по `ksrType` — здесь разбираем на плоский код->цена, категория
    дальше не нужна (принадлежность коду ресурса уже однозначна)."""
    data = json.loads(raw_json)
    prices: dict[str, float] = {}
    for group in data.get("items", []):
        for item in group.get("items", []):
            code = item.get("code")
            price = item.get("estimatedPrice")
            if code and price is not None:
                prices[code] = float(price)
    return prices


def parse_gosr_workbook(xlsx_bytes: bytes) -> dict[str, GosrIndexEntry]:
    """Оба листа отчёта (материалы/изделия/оборудование — лист 1, машины и
    механизмы — лист 2) имеют одинаковую структуру столбцов: код, название,
    единица, базисная цена на 01.01.2022, номер группы, название группы,
    индекс. Первые 3 строки — заголовок и подписи столбцов, данные с 4-й.

    Важно про машины: базисная цена на листе 2 — **без учёта оплаты труда
    машинистов** (см. заголовок листа), это только часть машино-часа,
    которую покрывает ГОСР-индекс. Зарплата машиниста требует отдельного
    источника (не подключён в этом фундаменте) — см. CLAUDE.md.
    """
    workbook = openpyxl.load_workbook(BytesIO(xlsx_bytes), data_only=True)
    entries: dict[str, GosrIndexEntry] = {}
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=4, values_only=True):
            code = row[0]
            if not code:
                continue
            entries[code] = GosrIndexEntry(
                resource_code=code,
                resource_name=row[1] or "",
                unit=row[2] or "",
                base_price_2022=float(row[3]) if row[3] is not None else 0.0,
                group_number=int(row[4]) if row[4] is not None else 0,
                group_name=row[5] or "",
                index_value=float(row[6]) if row[6] is not None else 1.0,
            )
    return entries


def build_regulatory_version(
    region_name: str,
    period_label: str,
    published_at: date,
    gosr_index: dict[str, GosrIndexEntry],
    resource_codes: Iterable[str],
) -> RegulatoryVersion:
    """Версия ГОСР-индексов для одного региона — через существующую модель
    Агента 10 (`RegulatoryVersion`/`RegulatoryUpdateStore.detect_update`/
    `approve_update`), не новая структура: смена квартала для одного и того
    же региона обязана пройти обычное обнаружение-diff-подтверждение
    эксперта, прежде чем Агент 4 станет использовать новые индексы.

    `resource_codes` — какие именно коды класть в версию: полный отчёт ГОСР
    на регион — десятки тысяч строк, а по факту нужны только те коды, что
    реально встречаются в каталоге ГЭСН, с которым сейчас работает Агент 4 —
    фильтруем явно, а не тащим всё целиком в каждую версию."""
    items = {
        code: str(gosr_index[code].index_value) for code in resource_codes if code in gosr_index
    }
    region_key = region_name.replace(" ", "_").replace(".", "").replace("г_", "")
    return RegulatoryVersion(
        source_id=GOSR_SOURCE_ID_TEMPLATE.format(region_key=region_key),
        version_label=period_label,
        published_at=published_at,
        items=items,
    )
