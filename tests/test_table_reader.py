"""Тесты извлечения СТРУКТУРЫ таблиц (не плоского текста) из .docx/.xlsx —
`document_analyst.table_reader`. Проверяют механику разбора (строки/ячейки
в порядке появления, несколько таблиц/листов), не конкретный реальный
документ — см. предупреждение в самом модуле про отсутствие проверки на
реальной ведомости объёмов работ."""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from docx import Document
from openpyxl import Workbook

from document_analyst import extract_docx_tables, extract_xlsx_tables


def _build_docx(tables: list[list[list[str]]]) -> bytes:
    document = Document()
    for table_rows in tables:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_idx, row in enumerate(table_rows):
            for col_idx, cell_text in enumerate(row):
                table.cell(row_idx, col_idx).text = cell_text
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _build_xlsx(sheets: list[list[list[str]]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for i, rows in enumerate(sheets):
        sheet = workbook.create_sheet(title=f"Лист{i + 1}")
        for row in rows:
            sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def test_extracts_single_docx_table_rows_in_order():
    content = _build_docx([[["Наименование работ", "Ед. изм.", "Количество"], ["Ремонт кровли", "м2", "250"]]])
    tables = extract_docx_tables(content)
    assert len(tables) == 1
    assert tables[0] == [["Наименование работ", "Ед. изм.", "Количество"], ["Ремонт кровли", "м2", "250"]]


def test_extracts_multiple_docx_tables_separately():
    content = _build_docx(
        [
            [["Заказчик", "ГБУЗ №1"]],
            [["Наименование работ", "Ед. изм.", "Количество"], ["Ремонт кровли", "м2", "250"]],
        ]
    )
    tables = extract_docx_tables(content)
    assert len(tables) == 2
    assert tables[0] == [["Заказчик", "ГБУЗ №1"]]
    assert tables[1][0] == ["Наименование работ", "Ед. изм.", "Количество"]


def test_docx_with_no_tables_returns_empty_list():
    document = Document()
    document.add_paragraph("Просто текст, без таблиц.")
    buf = io.BytesIO()
    document.save(buf)
    assert extract_docx_tables(buf.getvalue()) == []


def test_extracts_xlsx_sheet_rows_with_calculated_values():
    content = _build_xlsx(
        [[["Наименование работ", "Ед. изм.", "Количество"], ["Ремонт кровли", "м2", 250]]]
    )
    tables = extract_xlsx_tables(content)
    assert len(tables) == 1
    assert tables[0][0] == ["Наименование работ", "Ед. изм.", "Количество"]
    # openpyxl отдаёт число как значение ячейки — приводим к строке, как и
    # везде в этом читателе (см. докстринг: единый текстовый вид ячеек).
    assert tables[0][1] == ["Ремонт кровли", "м2", "250"]


def test_extracts_multiple_xlsx_sheets_separately():
    content = _build_xlsx(
        [
            [["Пояснительная записка"]],
            [["Наименование работ", "Ед. изм.", "Количество"], ["Ремонт кровли", "м2", 250]],
        ]
    )
    tables = extract_xlsx_tables(content)
    assert len(tables) == 2
    assert tables[0] == [["Пояснительная записка"]]
    assert tables[1][0] == ["Наименование работ", "Ед. изм.", "Количество"]


def test_xlsx_empty_rows_are_skipped_not_kept_as_blanks():
    content = _build_xlsx([[["Наименование работ"], [], ["Ремонт кровли"]]])
    tables = extract_xlsx_tables(content)
    assert tables[0] == [["Наименование работ"], ["Ремонт кровли"]]
