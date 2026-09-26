"""Тесты извлечения текста из .docx (`document_analyst.docx_reader`) —
проверяют механику (абзацы + таблицы, пустые абзацы отбрасываются), не
конкретный реальный документ (тот проверяется вручную, см. CLAUDE.md)."""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from docx import Document

from document_analyst import extract_text_from_docx


def _build_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_idx, row in enumerate(table_rows):
            for col_idx, cell_text in enumerate(row):
                table.cell(row_idx, col_idx).text = cell_text
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_extracts_paragraph_text_in_order():
    content = _build_docx(["Требования к участнику закупки.", "Опыт не менее 3 лет."])
    text = extract_text_from_docx(content)
    assert "Требования к участнику закупки." in text
    assert text.index("Требования") < text.index("Опыт не менее")


def test_skips_empty_paragraphs():
    content = _build_docx(["Первый абзац.", "", "   ", "Второй абзац."])
    text = extract_text_from_docx(content)
    lines = [line for line in text.split("\n") if line]
    assert lines == ["Первый абзац.", "Второй абзац."]


def test_includes_table_cell_text():
    content = _build_docx(
        ["Требования к содержанию заявки:"],
        table_rows=[["Раздел", "Опыт исполнения договора не менее 20 процентов НМЦК"]],
    )
    text = extract_text_from_docx(content)
    assert "Опыт исполнения договора не менее 20 процентов НМЦК" in text
