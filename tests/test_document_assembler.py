import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from docx import Document

from classifier.sample_tenders import SAMPLE_TENDERS
from document_analyst import extract_requirements
from document_analyst.models import ExtractedRequirements
from document_analyst.sample_documents import SAMPLE_DOCUMENTS
from document_assembler import assemble_document_package, generate_participant_info_docx
from onboarding import (
    Capacity,
    ClientProfile,
    CompletedContract,
    FinancialReadiness,
    LegalInfo,
    PermitsExperience,
    TaxRegimeChoice,
    validate_profile,
)


def make_ready_profile() -> ClientProfile:
    """Профиль, доведённый до статуса READY по полному пайплайну Агента 11 —
    без этого assemble_document_package() обязан отказать."""
    profile = ClientProfile(
        client_id="client-1",
        region_codes=["77"],
        legal=LegalInfo(
            org_name="ООО СтройМастер",
            inn="7701234567",
            ogrn="1027700132195",
            legal_address="г. Москва, ул. Примерная, д. 1",
            contact_person="Иванов Иван",
            phone="+79991234567",
            email="info@example.ru",
        ),
        permits_experience=PermitsExperience(
            sro_membership=True,
            sro_number="СРО-С-123-456",
            completed_contracts=[
                CompletedContract(
                    object_name="Капремонт школы №5", customer="ДепОбр", amount=5_000_000, year=2024
                )
            ],
            years_of_experience=5,
        ),
        capacity=Capacity(staff_count=15, own_workforce_description="15 штатных рабочих"),
        financial=FinancialReadiness(
            tax_regime=TaxRegimeChoice.USN_6_NO_VAT,
            avg_annual_revenue=50_000_000,
            working_capital=3_000_000,
            bank_guarantee_available=True,
        ),
    )
    validate_profile(profile)
    profile.submit_expert_review(reviewer="Edwin", approved=True)
    profile.mark_ready()
    return profile


def find_tender(purchase_number: str):
    return next(t for t in SAMPLE_TENDERS if t.purchase_number == purchase_number)


def test_rejects_profile_that_is_not_ready():
    profile = ClientProfile(client_id="client-2")  # остаётся DRAFT
    tender = find_tender("0173200001426000101")

    try:
        assemble_document_package(profile, tender)
        assert False, "ожидался ValueError"
    except ValueError as exc:
        assert "READY" in str(exc)
        assert "DRAFT" in str(exc)


def test_assembles_package_skeleton_for_ready_profile():
    profile = make_ready_profile()
    tender = find_tender("0173200001426000101")  # капремонт кровли, требует СРО и опыт

    package = assemble_document_package(profile, tender)

    assert package.client_id == "client-1"
    assert package.tender_purchase_number == tender.purchase_number
    assert package.is_complete()
    assert package.missing_fields() == []

    by_name = {f.name: f for f in package.fields}
    assert by_name["ИНН"].value == "7701234567"
    assert by_name["ИНН"].source == "legal.inn"
    assert by_name["ИНН"].status == "ready"
    assert by_name["Членство в СРО (номер)"].status == "ready"
    assert by_name["Номер закупки"].source == "tender.purchase_number"
    # Значение поля — человекочитаемая подпись выбранной комбинации
    # режим+ставка (агент 5 берёт готовую ставку прямо из этого выбора,
    # не вычисляет её из дохода), а не голое имя enum-константы.
    assert by_name["Система налогообложения"].value == "УСН 6% без НДС"
    assert by_name["Система налогообложения"].source == "financial.tax_regime"


def test_sro_field_not_applicable_when_tender_does_not_require_it():
    profile = make_ready_profile()
    profile.permits_experience.sro_number = ""  # у клиента нет номера СРО
    tender = find_tender("0123200004426000303")  # штукатурные работы, СРО не требуется

    package = assemble_document_package(profile, tender)

    by_name = {f.name: f for f in package.fields}
    assert by_name["Членство в СРО (номер)"].status == "not_applicable"
    assert package.is_complete()  # not_applicable не считается пропущенным полем


def test_missing_required_field_is_flagged():
    profile = make_ready_profile()
    profile.legal.contact_person = ""
    tender = find_tender("0173200001426000101")

    # profile уже READY (статус не пересчитывается при прямом изменении поля
    # после mark_ready — это ожидаемо: правки после READY должны идти через
    # новый цикл проверки, тут просто проверяем, что сборщик увидит дыру)
    package = assemble_document_package(profile, tender)

    by_name = {f.name: f for f in package.fields}
    assert by_name["Контактное лицо"].status == "missing"
    assert not package.is_complete()
    assert any(f.name == "Контактное лицо" for f in package.missing_fields())


def _reviewed_requirements(purchase_number: str) -> ExtractedRequirements:
    """Реальный прогон Агента 3 на тестовом документе + подтверждение
    эксперта — подключение к Агенту 6 не отменяет эту проверку."""
    extracted = extract_requirements(purchase_number, SAMPLE_DOCUMENTS[purchase_number])
    extracted.mark_expert_reviewed(reviewer="Edwin")
    return extracted


def test_rejects_unreviewed_agent_3_output():
    """Подключение Агента 3 к Агенту 6 не автоматизирует обязательную
    проверку эксперта из протокола контроля качества."""
    profile = make_ready_profile()
    tender = find_tender("0173200001426000101")
    extracted = extract_requirements(tender.purchase_number, SAMPLE_DOCUMENTS[tender.purchase_number])
    assert not extracted.expert_reviewed

    try:
        assemble_document_package(profile, tender, extracted_requirements=extracted)
        assert False, "ожидался ValueError"
    except ValueError as exc:
        assert "проверен экспертом" in str(exc)


def test_rejects_agent_3_output_for_a_different_tender():
    profile = make_ready_profile()
    tender = find_tender("0173200001426000101")
    other_purchase_number = "0350200003426000202"
    extracted = _reviewed_requirements(other_purchase_number)

    try:
        assemble_document_package(profile, tender, extracted_requirements=extracted)
        assert False, "ожидался ValueError"
    except ValueError as exc:
        assert tender.purchase_number in str(exc)
        assert other_purchase_number in str(exc)


def test_agent_3_contract_security_finding_makes_field_required():
    """Tender вообще не хранит требование к обеспечению исполнения
    контракта — этот факт есть только в тексте документации, который
    разбирает Агент 3. Без него поле не появляется в пакете вовсе; с ним —
    появляется и становится обязательным, если у клиента нет банковской
    гарантии."""
    profile = make_ready_profile()
    profile.financial.bank_guarantee_available = False
    tender = find_tender("0173200001426000101")  # документ содержит обеспечение контракта 10%

    package_without_agent_3 = assemble_document_package(profile, tender)
    assert "Обеспечение исполнения контракта (банковская гарантия)" not in {
        f.name for f in package_without_agent_3.fields
    }

    extracted = _reviewed_requirements(tender.purchase_number)
    package = assemble_document_package(profile, tender, extracted_requirements=extracted)

    by_name = {f.name: f for f in package.fields}
    contract_security = by_name["Обеспечение исполнения контракта (банковская гарантия)"]
    assert contract_security.required is True
    assert contract_security.status == "missing"
    assert not package.is_complete()


def test_agent_3_participant_finding_overrides_tender_when_tender_says_not_required():
    """Документация может требовать СРО, даже если структурированные данные
    закупки (Tender.requires_sro) этого не показывают — например, если
    данные закупки устарели или неполны. Агент 3 — независимая сверка."""
    profile = make_ready_profile()
    profile.permits_experience.sro_number = ""
    tender = find_tender("0123200004426000303")  # requires_sro=False у самого Tender

    extracted = extract_requirements(
        tender.purchase_number,
        SAMPLE_DOCUMENTS[tender.purchase_number] + "\nУчастник должен быть членом СРО.\n",
    )
    extracted.mark_expert_reviewed(reviewer="Edwin")

    package = assemble_document_package(profile, tender, extracted_requirements=extracted)

    by_name = {f.name: f for f in package.fields}
    assert by_name["Членство в СРО (номер)"].required is True
    assert by_name["Членство в СРО (номер)"].status == "missing"


def test_agent_3_hidden_risks_flow_into_package_without_blocking_completeness():
    profile = make_ready_profile()
    tender = find_tender("0173200001426000101")  # документ содержит два скрытых риска
    extracted = _reviewed_requirements(tender.purchase_number)

    package = assemble_document_package(profile, tender, extracted_requirements=extracted)

    assert len(package.hidden_risks) == 2
    assert package.hidden_risks == extracted.hidden_risks
    # Риски — не часть fields (проверку комплектности не видят), поэтому
    # они физически не могут её заблокировать.
    assert not any("риск" in f.name.lower() for f in package.fields)
    assert package.is_complete()


# -- Генерация черновика п.1/п.5 ст.43 44-ФЗ (2026-09-27, CLAUDE.md, п.2) ---


def _docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_package_includes_generated_participant_info_docx():
    profile = make_ready_profile()
    tender = find_tender("0173200001426000101")

    package = assemble_document_package(profile, tender)

    assert package.participant_info_docx is not None
    text = _docx_text(package.participant_info_docx)
    assert profile.legal.org_name in text
    assert profile.legal.inn in text
    assert tender.purchase_number in text
    # Единой типовой формы не существует — документ должен честно говорить
    # об этом, не выглядеть как заявка "по форме".
    assert "не заявка по единой форме" in text.lower()


def test_generated_docx_marks_smp_declaration_as_needing_clarification_when_unset():
    profile = make_ready_profile()
    assert profile.legal.is_smp is None  # по умолчанию, не заполнено

    content = generate_participant_info_docx(profile, find_tender("0173200001426000101"))

    text = _docx_text(content)
    assert "ТРЕБУЕТ УТОЧНЕНИЯ" in text


def test_generated_docx_includes_smp_declaration_when_profile_confirms_it():
    profile = make_ready_profile()
    profile.legal.is_smp = True

    content = generate_participant_info_docx(profile, find_tender("0173200001426000101"))

    text = _docx_text(content)
    assert "субъектам малого предпринимательства" in text.lower()
    assert "ТРЕБУЕТ УТОЧНЕНИЯ" not in text


def test_generated_docx_states_no_smp_declaration_needed_when_profile_denies_it():
    profile = make_ready_profile()
    profile.legal.is_smp = False

    content = generate_participant_info_docx(profile, find_tender("0173200001426000101"))

    text = _docx_text(content)
    assert "не относится к субъектам малого предпринимательства" in text.lower()
    assert "ТРЕБУЕТ УТОЧНЕНИЯ" not in text


def test_package_lists_manual_sections_for_object_proposal_and_price():
    """П.2 (предложение по объекту закупки) и п.3 (предложение о цене) ст.43
    44-ФЗ — сознательно не генерируются никогда, независимо от профиля/
    закупки (см. CLAUDE.md, открытый п.2)."""
    profile = make_ready_profile()
    tender = find_tender("0173200001426000101")

    package = assemble_document_package(profile, tender)

    assert len(package.manual_sections) == 2
    names = {s.name for s in package.manual_sections}
    assert "Предложение по объекту закупки" in names
    assert "Предложение о цене контракта" in names
    for section in package.manual_sections:
        assert "ст. 43" in section.article_reference
        assert section.note  # объяснение не пустое

    # Часть 1/Часть 2 — устаревшая терминология (ст.66 утратила силу
    # 01.01.2022), не должна просочиться в новый код.
    assert not any("часть 1" in s.note.lower() or "часть 2" in s.note.lower() for s in package.manual_sections)


def test_manual_sections_do_not_affect_completeness_check():
    """manual_sections — не часть fields, поэтому Агент 7 (комплектность)
    их не видит и они не блокируют is_complete()."""
    profile = make_ready_profile()
    tender = find_tender("0173200001426000101")

    package = assemble_document_package(profile, tender)

    assert package.manual_sections  # непустой список
    assert package.is_complete()  # но не мешает считать пакет "полным" по fields
